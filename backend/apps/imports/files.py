"""Bounded passive tabular parsing. No workbook execution or permanent files."""
import csv
import io
import zipfile

from defusedxml import ElementTree
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 40
MAX_CELL = 4000


def reject():
    raise ValidationError({'file': 'Xavfsiz UTF-8 CSV yoki oddiy XLSX fayl tanlang (2 MB, 1000 qator, 40 ustungacha).'})


def read_table(upload):
    if upload.size > MAX_BYTES:
        reject()
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        reject()
    name = upload.name.lower()
    try:
        if name.endswith('.csv'):
            if b'\x00' in raw or raw.startswith(b'PK'):
                reject()
            table = csv.reader(io.StringIO(raw.decode('utf-8-sig')), strict=True)
            values = []
            for row in table:
                values.append(row)
                if len(values) > MAX_ROWS + 1:
                    reject()
        elif name.endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 100 or sum(i.file_size for i in entries) > 8 * MAX_BYTES:
                    reject()
                for entry in entries:
                    lower = entry.filename.lower()
                    if (entry.flag_bits & 1 or entry.file_size > 4 * MAX_BYTES
                        or entry.file_size > max(entry.compress_size, 1) * 100
                        or any(x in lower for x in ('vba', 'externallink', 'embeddings', 'connections', 'querytable', 'activex'))):
                        reject()
                    if lower.endswith(('.xml', '.rels')):
                        data = archive.read(entry)
                        root = ElementTree.fromstring(data)
                        for element in root.iter():
                            if element.attrib.get('TargetMode') == 'External':
                                reject()
                if '[Content_Types].xml' not in archive.namelist():
                    reject()
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    reject()
                sheet = book.worksheets[0]
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                    reject()
                values = []
                for row in sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1):
                    populated = [cell for cell in row if cell.value is not None]
                    if not populated:
                        continue
                    if len(values) >= MAX_ROWS + 1 or any(c.data_type == 'f' for c in populated):
                        reject()
                    last = max(c.column for c in populated)
                    values.append(['' if c.value is None else str(c.value) for c in row[:last]])
            finally:
                book.close()
        else:
            reject()
        if len(values) < 2:
            reject()
        width = len(values[0])
        if width > MAX_COLUMNS or not width or any(not h.strip() for h in values[0]):
            reject()
        headers = [h.strip() for h in values[0]]
        if len(set(headers)) != width:
            reject()
        rows = []
        for row in values[1:]:
            if not row or not any(str(v).strip() for v in row):
                continue
            if len(row) > width:
                reject()
            row = [str(v).strip() for v in row] + [''] * (width - len(row))
            if any(len(v) > MAX_CELL or v.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) for v in row):
                reject()
            rows.append(row)
        if not rows or len(rows) > MAX_ROWS:
            reject()
        return headers, rows
    except ValidationError:
        raise
    except Exception:
        reject()
