import io
from pathlib import Path
from xml.sax.saxutils import escape

from django.conf import settings
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def document_data(plan):
    if plan.snapshot:
        return plan.snapshot
    from apps.plans.serializers import ContinuityPlanSerializer

    return {
        **dict(ContinuityPlanSerializer(plan).data),
        "organization_name": plan.organization.name,
        "contacts": plan.contacts,
        "approved_by_username": plan.approved_by.username if plan.approved_by else None,
    }


def sections(data):
    return [
        (label, str(data.get(key) or "Kiritilmagan"))
        for key, label in [
            ("scope", "Maqsad va qamrov"),
            ("activation_conditions", "Ishga tushirish shartlari"),
            ("contacts", "Mas’ullar va aloqalar"),
            ("temporary_procedures", "Vaqtinchalik ish tartibi"),
            ("return_to_normal", "Normal holatga qaytish"),
            ("verification_criteria", "Tekshirish mezonlari"),
        ]
    ]


def export_pdf(data):
    font = Path(settings.BASE_DIR) / "apps/plans/fonts/DejaVuSerif.ttf"
    if "BCMS" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("BCMS", str(font)))
    body = ParagraphStyle(
        "UzbekBody", fontName="BCMS", fontSize=10, leading=15, spaceAfter=8, alignment=TA_LEFT
    )
    heading = ParagraphStyle(
        "UzbekHeading",
        parent=body,
        fontSize=14,
        leading=19,
        textColor=colors.HexColor("#12395d"),
        spaceBefore=12,
    )

    def p(value, style=body):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)

    buffer = io.BytesIO()
    story = [
        p("RiskGuard · BCMS Designer", heading),
        p(data["title"], heading),
        p(f"{data.get('organization_name', '')} · Versiya {data['version']} · {data['status']}"),
        p(f"Jarayon: {data.get('process_name', '')}"),
        p(
            f"Tasdiqlovchi: {data.get('approved_by_username') or 'Tasdiqlanmagan'} · {data.get('approved_at') or ''}"
        ),
        Spacer(1, 12),
    ]
    for title, content in sections(data):
        story.extend([p(title, heading), p(content)])
    story.append(p("Tiklash bosqichlari", heading))
    rows = [[p("№"), p("Bosqich va tavsif"), p("Mas’ul"), p("Daqiqa")]]
    for step in data.get("steps", []):
        rows.append(
            [
                p(step["sequence"]),
                p(str(step["title"]) + "\n" + str(step.get("description", ""))),
                p(step.get("assigned_name") or "Belgilanmagan"),
                p(step["expected_duration_minutes"]),
            ]
        )
    table = Table(rows, colWidths=[25, 270, 120, 65], repeatRows=1, hAlign="LEFT", splitInRow=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5f3f1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(table)
    schedule = data.get("schedule_snapshot")
    if schedule:
        story.append(p("Saqlangan tiklash jadvali va resurslar", heading))
        for row in schedule.get("rows", []):
            story.append(
                p(
                    f"{row['process_name']}: {row['start_minutes']}–{row['end_minutes']} daqiqa; {row['cost']} UZS; resurslar: {row['assigned_resources']}"
                )
            )

    def footer(canvas, doc):
        if doc.page == 1:
            canvas.saveState()
            canvas.setStrokeColor(colors.HexColor('#c6b48a'))
            canvas.setLineWidth(0.4)
            canvas.translate(doc.pagesize[0]-55, doc.pagesize[1]-55)
            canvas.rotate(45)
            for size in [20,35,50]:
                canvas.rect(-size,-size,size*2,size*2,fill=0,stroke=1)
            canvas.restoreState()
        canvas.setFont("BCMS", 8)
        canvas.drawString(48, 28, f"RiskGuard · {doc.page} · Sertifikatlash hujjati emas")

    SimpleDocTemplate(buffer, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=45).build(
        story, onFirstPage=footer, onLaterPages=footer
    )
    return buffer.getvalue()


def export_docx(data):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Serif", Pt(11)
    doc.add_heading("RiskGuard · BCMS Designer", 0)
    doc.add_heading(data["title"], 1)
    doc.add_paragraph(
        f"{data.get('organization_name', '')} · Versiya {data['version']} · {data['status']}"
    )
    doc.add_paragraph(f"Jarayon: {data.get('process_name', '')}")
    doc.add_paragraph(
        f"Tasdiqlovchi: {data.get('approved_by_username') or 'Tasdiqlanmagan'} · {data.get('approved_at') or ''}"
    )
    for label, content in sections(data):
        doc.add_heading(label, 2)
        doc.add_paragraph(content)
    doc.add_heading("Tiklash bosqichlari", 2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["№", "Bosqich va tavsif", "Mas’ul", "Daqiqa"]):
        cell.text = label
    for step in data.get("steps", []):
        for cell, value in zip(
            table.add_row().cells,
            [
                step["sequence"],
                step["title"] + "\n" + step.get("description", ""),
                step.get("assigned_name") or "Belgilanmagan",
                step["expected_duration_minutes"],
            ],
        ):
            cell.text = str(value)
    if data.get("schedule_snapshot"):
        doc.add_heading("Saqlangan jadval va resurslar", 2)
        for row in data["schedule_snapshot"].get("rows", []):
            doc.add_paragraph(
                f"{row['process_name']}: {row['start_minutes']}–{row['end_minutes']} daqiqa; {row['cost']} UZS; {row['assigned_resources']}"
            )
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
