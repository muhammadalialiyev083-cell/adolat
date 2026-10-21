"""Real database acceptance journey; tenant, RBAC, validation and scheduler checks."""

import io
import json
from decimal import Decimal
from unittest.mock import patch

import pytest
from apps.authentication.models import User
from apps.bia.models import BIAAssessment
from apps.core.models import AuditLog
from apps.engine.scheduler import compare_schedules, schedule
from apps.organizations.models import Organization, Service
from apps.plans.models import ContinuityPlan
from apps.processes.models import Asset, BusinessProcess
from django.core.cache import cache
from docx import Document
from pypdf import PdfReader
from rest_framework.test import APIClient

PASSWORD = "Workflow-Example!71264"


@pytest.fixture(autouse=True)
def independent_cache():
    cache.clear()
    yield
    cache.clear()


@pytest.fixture
def data(db):
    org = Organization.objects.create(name="Asosiy universitet", sector="Ta’lim")
    other = Organization.objects.create(name="Boshqa tashkilot", sector="Ta’lim")
    admin = User.objects.create_user(
        username="org_admin", password=PASSWORD, role="ADMIN", organization=org
    )
    owner = User.objects.create_user(username="owner", role="PROCESS_OWNER", organization=org)
    auditor = User.objects.create_user(username="auditor", role="AUDITOR", organization=org)
    service = Service.objects.create(organization=org, name="Xizmat", owner=admin)
    process = BusinessProcess.objects.create(
        organization=org, service=service, name="Jarayon", owner=owner
    )
    other_user = User.objects.create_user(username="foreign", organization=other)
    other_service = Service.objects.create(organization=other, name="Yashirin xizmat")
    other_process = BusinessProcess.objects.create(
        organization=other, service=other_service, name="Yashirin jarayon"
    )
    asset = Asset.objects.create(organization=org, name="Server", type="SERVER")
    other_asset = Asset.objects.create(organization=other, name="Maxfiy server", type="SERVER")
    client = APIClient()
    client.force_authenticate(admin)
    return locals()


def create(client, resource, values):
    response = client.post("/api/v1/" + resource + "/", values, format="json")
    assert response.status_code == 201, response.data
    return response.data


def plan_values(process, **extra):
    return {
        "process": process,
        "title": "O‘zbekcha o‘quv rejasi",
        "version": "1.0",
        "scope": "Raqamli ta’lim",
        "activation_conditions": "Xizmat uzilganda mas’ul hodisani tasdiqlaydi.",
        "contacts": "Mas’ul: test@example.invalid",
        "temporary_procedures": "Oflayn ro‘yxat.",
        "return_to_normal": "Ma’lumotlar solishtiriladi.",
        "verification_criteria": "Nazorat so‘rovi va ma’lumot yaxlitligi.",
        "steps": [
            {
                "sequence": 1,
                "title": "Zaxira serverni tekshirish",
                "description": "O‘quv tartibi",
                "expected_duration_minutes": 30,
            }
        ],
        **extra,
    }


@pytest.mark.django_db
def test_full_csrf_journey_login_organization_to_export_logout():
    initial = User.objects.create_superuser(username="initial", password=PASSWORD, role="ADMIN")
    client = APIClient(enforce_csrf_checks=True)
    csrf = client.get("/api/v1/auth/csrf").data["csrfToken"]
    response = client.post(
        "/api/v1/auth/login",
        {"username": "initial", "password": PASSWORD},
        format="json",
        HTTP_X_CSRFTOKEN=csrf,
    )
    assert response.status_code == 200
    client.credentials(HTTP_X_CSRFTOKEN=response.data["csrfToken"])
    org = create(
        client,
        "organizations",
        {"name": "API universiteti", "sector": "Ta’lim", "contact_details": "test@example.invalid"},
    )
    dept = create(client, "departments", {"name": "Raqamlashtirish"})
    service = create(client, "services", {"name": "Ta’lim", "owner": initial.pk})
    proc = create(
        client,
        "processes",
        {
            "name": "Ta’lim jarayoni",
            "description": "API tekshiruvi",
            "department": dept["id"],
            "service": service["id"],
            "owner": initial.pk,
            "criticality": "CRITICAL",
        },
    )
    asset = create(client, "assets", {"name": "Ta’lim serveri", "type": "SERVER"})
    create(
        client,
        "dependencies",
        {"from_type": "PROCESS", "from_id": proc["id"], "to_type": "ASSET", "to_id": asset["id"]},
    )
    bia = create(
        client,
        "bia",
        {
            "process": proc["id"],
            "rto_hours": "4.00",
            "mtpd_hours": "12.00",
            "rpo_minutes": 30,
            "backup_interval_minutes": 60,
            "rationale": "Ekspert bahosi",
            "data_source": "O‘quv ma’lumoti",
            "impacts": [
                {
                    "horizon_hours": "4.00",
                    "financial_loss": "1000000",
                    "operational_score": 4,
                    "legal_score": 3,
                    "reputational_score": 2,
                }
            ],
        },
    )
    risk = create(
        client,
        "risks",
        {
            "process": proc["id"],
            "asset": asset["id"],
            "threat": "Server nosozligi",
            "likelihood": 4,
            "impact": 5,
        },
    )
    assert risk["risk_score"] == 20
    create(
        client, "controls", {"risk": risk["id"], "title": "Backup tekshiruvi", "owner": initial.pk}
    )
    resource = create(client, "resources", {"name": "Uskuna", "capacity": 1})
    create(
        client,
        "strategies",
        {
            "process": proc["id"],
            "name": "Backup",
            "estimated_rto_hours": "1.00",
            "estimated_rpo_minutes": 15,
            "estimated_cost": "500000",
            "resource_requirements": {str(resource["id"]): 1},
        },
    )
    run = create(client, "schedule-runs", {"staff_capacity": 1, "budget": "1000000"})
    assert run["results"]["constrained"]["feasible"] is True
    assert run["results"]["constrained"]["rows"][0]["end_minutes"] == 60
    assert run["input_snapshot"]["tasks"][0]["target"] == 240
    plan = create(client, "plans", plan_values(proc["id"], schedule_run=run["id"]))
    assert client.post(f"/api/v1/bia/{bia['id']}/approve/").status_code == 200
    approved = client.post(f"/api/v1/plans/{plan['id']}/publish/")
    assert approved.status_code == 200, approved.data
    snapshot = approved.data["plan"]["snapshot"]
    assert snapshot["organization_name"] == org["name"]
    assert (
        client.patch(f"/api/v1/plans/{plan['id']}/", {"title": "Unsafe"}, format="json").status_code
        == 400
    )
    assert client.delete(f"/api/v1/plans/{plan['id']}/").status_code == 400
    revision = client.post(f"/api/v1/plans/{plan['id']}/revise/")
    assert revision.status_code == 201, revision.data
    assert revision.data["previous_version"] == plan["id"] and revision.data["status"] == "DRAFT"
    assert (
        client.patch(
            f"/api/v1/plans/{revision.data['id']}/", {"title": "Yangi versiya"}, format="json"
        ).status_code
        == 200
    )
    assert ContinuityPlan.objects.get(pk=plan["id"]).snapshot == snapshot
    assert (
        client.patch(f"/api/v1/bia/{bia['id']}/", {"rto_hours": 2}, format="json").status_code
        == 400
    )
    revised_bia = client.post(f"/api/v1/bia/{bia['id']}/revise/")
    assert revised_bia.status_code == 201, revised_bia.data
    assert (
        client.patch(
            f"/api/v1/bia/{revised_bia.data['id']}/", {"rto_hours": 3}, format="json"
        ).status_code
        == 200
    )
    scenario = create(
        client,
        "scenarios",
        {
            "title": "Server uzilishi",
            "disrupted_asset_ids": [asset["id"]],
            "estimated_outage_hours": "8.00",
        },
    )
    simulation = client.post(f"/api/v1/scenarios/{scenario['id']}/simulate/")
    assert simulation.status_code == 200 and simulation.data["affected_processes_count"] == 1
    exercise = create(
        client,
        "exercises",
        {
            "plan": plan["id"],
            "scenario": scenario["id"],
            "exercise_date": "2026-10-21",
            "actual_recovery_hours": "2.00",
            "findings": "Aloqa ro‘yxati eskirgan",
            "improvement_actions": "Kontaktni yangilash",
        },
    )
    assert exercise["planned_recovery_hours"] == "4.00"
    report = client.get("/api/v1/reports/organization").data
    assert (
        report["summary"]["rto_compliance_percent"] == 100.0
        and report["summary"]["open_findings_count"] == 1
    )
    for ext in ["pdf", "docx"]:
        response = client.get(f"/api/v1/plans/{plan['id']}/export_{ext}/")
        assert response.status_code == 200
        text = (
            "".join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
            if ext == "pdf"
            else " ".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
        )
        assert (
            "O‘zbekcha o‘quv rejasi" in text and org["name"] in text and "Yangi versiya" not in text
        )
    assert client.post(f"/api/v1/plans/{plan['id']}/archive/").status_code == 200
    assert ContinuityPlan.objects.get(pk=plan["id"]).snapshot == snapshot
    assert client.get("/api/v1/reports/processes.csv").status_code == 200
    assert client.get("/api/v1/processes/").data["results"][0]["name"] == proc["name"]
    assert client.get("/api/v1/dependency-graph").data["nodes"]
    assert client.get("/api/v1/workflow").data["completed"] >= 6
    assert PASSWORD not in json.dumps(list(AuditLog.objects.values("before_state", "after_state")))
    changed = client.post(
        "/api/v1/auth/password",
        {"old_password": PASSWORD, "new_password": "Changed-Workflow!73589"},
        format="json",
    )
    assert changed.status_code == 200, changed.data
    client.credentials(HTTP_X_CSRFTOKEN=changed.data["csrfToken"])
    assert client.get("/api/v1/auth/me").status_code == 200
    assert client.post("/api/v1/auth/logout").status_code == 200
    assert client.get("/api/v1/auth/me").status_code == 403


@pytest.mark.django_db
@pytest.mark.parametrize(
    "resource,values",
    [
        ("processes", {"name": "Attack", "service": "other_service"}),
        ("assets", {"name": "Attack", "type": "SERVER", "owner": "other_user"}),
        ("bia", {"process": "other_process", "rto_hours": 1, "mtpd_hours": 2}),
        ("risks", {"process": "process", "asset": "other_asset", "threat": "Attack"}),
        (
            "strategies",
            {
                "process": "other_process",
                "name": "Attack",
                "estimated_rto_hours": 1,
                "estimated_rpo_minutes": 1,
            },
        ),
        (
            "dependencies",
            {
                "from_type": "PROCESS",
                "from_id": "process",
                "to_type": "ASSET",
                "to_id": "other_asset",
            },
        ),
        ("scenarios", {"title": "Attack", "disrupted_asset_ids": ["other_asset"]}),
    ],
)
def test_foreign_key_tenant_boundary(data, resource, values):
    payload = {
        k: [data[v[0]].pk]
        if isinstance(v, list)
        else data[v].pk
        if isinstance(v, str) and v in data and hasattr(data[v], "pk")
        else v
        for k, v in values.items()
    }
    response = data["client"].post("/api/v1/" + resource + "/", payload, format="json")
    assert response.status_code == 400, response.data


@pytest.mark.django_db
def test_cross_org_reads_search_exports_even_for_superuser(data):
    client = data["client"]
    data["admin"].is_superuser = True
    data["admin"].save()
    assert client.get(f"/api/v1/processes/{data['other_process'].pk}/").status_code == 404
    assert (
        client.patch(
            f"/api/v1/processes/{data['other_process'].pk}/", {"name": "attack"}, format="json"
        ).status_code
        == 404
    )
    assert client.get("/api/v1/processes/?search=Yashirin").data["count"] == 0
    assert "Yashirin" not in client.get("/api/v1/reports/processes.csv").content.decode("utf-8")


@pytest.mark.django_db
@pytest.mark.parametrize(
    "role", ["PROCESS_OWNER", "BCMS_MANAGER", "RISK_ANALYST", "APPROVER", "AUDITOR"]
)
def test_admin_only_user_management(data, role):
    user = User.objects.create_user(username="role_" + role, role=role, organization=data["org"])
    client = data["client"]
    client.force_authenticate(user)
    assert (
        client.post(
            "/api/v1/users/",
            {"username": "intruder", "password": PASSWORD, "role": "ADMIN"},
            format="json",
        ).status_code
        == 403
    )
    assert (
        client.patch(f"/api/v1/users/{user.pk}/", {"role": "ADMIN"}, format="json").status_code
        == 403
    )
    user.refresh_from_db()
    assert user.role == role


@pytest.mark.django_db
def test_process_owner_cannot_reassign_or_write_others(data):
    client = data["client"]
    client.force_authenticate(data["owner"])
    assert (
        client.patch(
            f"/api/v1/processes/{data['process'].pk}/", {"owner": data["admin"].pk}, format="json"
        ).status_code
        == 400
    )
    proc = BusinessProcess.objects.create(
        organization=data["org"], service=data["service"], name="Boshqa mas’ul", owner=data["admin"]
    )
    bia = BIAAssessment.objects.create(
        organization=data["org"], process=proc, rto_hours=1, mtpd_hours=2
    )
    assert client.post(f"/api/v1/bia/{bia.pk}/revise/").status_code == 403
    assert (
        client.patch(
            f"/api/v1/processes/{data['process'].pk}/",
            {"organization": data["other"].pk},
            format="json",
        ).status_code
        == 200
    )
    data["process"].refresh_from_db()
    assert data["process"].organization_id == data["org"].pk


@pytest.mark.django_db
@pytest.mark.parametrize(
    "values",
    [
        {"rto_hours": 0, "mtpd_hours": 4},
        {"rto_hours": -1, "mtpd_hours": 4},
        {"rto_hours": 4, "mtpd_hours": 4},
        {"rto_hours": 1, "mtpd_hours": 4, "impact_weights": {"financial": 2}},
        {
            "rto_hours": 1,
            "mtpd_hours": 4,
            "impacts": [{"horizon_hours": 1, "operational_score": 6}],
        },
        {"rto_hours": 1, "mtpd_hours": 4, "impacts": [{"horizon_hours": 1, "financial_loss": -1}]},
        {"rto_hours": 1, "mtpd_hours": 4, "impacts": [{"horizon_hours": 1}, {"horizon_hours": 1}]},
        {"rto_hours": "NaN", "mtpd_hours": 4},
    ],
)
def test_bia_boundary_validation(data, values):
    before = BIAAssessment.objects.count()
    response = data["client"].post(
        "/api/v1/bia/", {"process": data["process"].pk, **values}, format="json"
    )
    assert response.status_code == 400, response.data
    assert BIAAssessment.objects.count() == before


@pytest.mark.django_db
def test_nested_updates_are_atomic(data):
    client = data["client"]
    bia = create(
        client,
        "bia",
        {
            "process": data["process"].pk,
            "rto_hours": 1,
            "mtpd_hours": 4,
            "impacts": [{"horizon_hours": 1, "financial_loss": 10}],
        },
    )
    url = f"/api/v1/bia/{bia['id']}/"
    response = client.patch(
        url, {"impacts": [{"horizon_hours": 2, "financial_loss": 20}]}, format="json"
    )
    assert response.status_code == 200, response.data
    assert Decimal(response.data["impacts"][0]["horizon_hours"]) == 2
    with patch(
        "apps.bia.serializers.BIAImpact.objects.create", side_effect=RuntimeError("rollback test")
    ):
        with pytest.raises(RuntimeError):
            client.patch(
                url,
                {"comments": "Should rollback", "impacts": [{"horizon_hours": 3}]},
                format="json",
            )
    saved = BIAAssessment.objects.get(pk=bia["id"])
    assert saved.comments == "" and saved.impacts.first().horizon_hours == 2


@pytest.mark.django_db
@pytest.mark.parametrize(
    "payload",
    [
        {"likelihood": 0, "impact": 4},
        {"likelihood": "broken", "impact": 4},
        {"likelihood": 6, "impact": 5},
        {"likelihood": 4, "impact": 4, "backup_interval_minutes": -1},
    ],
)
def test_bad_engine_input_is_400(data, payload):
    assert (
        data["client"].post("/api/v1/analyses/evaluate-risk", payload, format="json").status_code
        == 400
    )


@pytest.mark.django_db
def test_approval_bypass_delete_protection_and_formula_csv(data):
    c = data["client"]
    p = data["process"].pk
    assert (
        c.post(
            "/api/v1/bia/",
            {"process": p, "rto_hours": 1, "mtpd_hours": 2, "status": "APPROVED"},
            format="json",
        ).status_code
        == 400
    )
    plan = create(c, "plans", {"process": p, "title": "Incomplete"})
    assert c.post(f"/api/v1/plans/{plan['id']}/publish/").status_code == 400
    create(
        c,
        "dependencies",
        {"from_type": "PROCESS", "from_id": p, "to_type": "ASSET", "to_id": data["asset"].pk},
    )
    assert c.delete(f"/api/v1/assets/{data['asset'].pk}/").status_code == 409
    assert c.delete(f"/api/v1/services/{data['service'].pk}/").status_code == 409
    data["process"].name = '=HYPERLINK("test")'
    data["process"].save()
    assert "'=HYPERLINK" in c.get("/api/v1/reports/processes.csv").content.decode("utf-8")
    unused = create(c, "assets", {"name": "Unused", "type": "NETWORK"})
    assert c.delete(f"/api/v1/assets/{unused['id']}/").status_code == 204
    assert c.get(f"/api/v1/assets/{unused['id']}/").status_code == 404
    c.force_authenticate(data["auditor"])
    assert (
        c.post("/api/v1/assets/", {"name": "attack", "type": "SERVER"}, format="json").status_code
        == 403
    )
    assert c.get("/api/v1/reports/organization").data["summary"]["rto_compliance_percent"] is None
    result = c.post("/api/v1/analyses/priority", {}, format="json").data["ranked_processes"][0]
    assert result["priority_score"] is None and result["status"] == "INSUFFICIENT_DATA"


def example():
    return {
        "tasks": [
            {
                "id": i,
                "name": str(i),
                "criticality": "CRITICAL" if i == 1 else "LOW",
                "priority": 0.8 if i == 1 else 0.2,
                "target": 180,
                "rpo": 30,
                "options": [
                    {
                        "id": i,
                        "name": "Backup",
                        "duration": 60,
                        "preparation": 0,
                        "staff": 1,
                        "resources": {"r": 1},
                        "rpo": 15,
                        "cost": "100",
                        "conditions": "",
                    }
                ],
            }
            for i in [1, 2]
        ],
        "resources": {
            "r": {"name": "Server", "capacity": 2, "available": 0},
            "staff": {"name": "Staff", "capacity": 2, "available": 0},
        },
        "dependencies": [{"from_type": "PROCESS", "from_id": 1, "to_type": "PROCESS", "to_id": 2}],
        "budget": "200",
    }


def test_hand_computed_dependency_and_utilization():
    results = compare_schedules(example())
    a = results["baseline"]
    b = results["constrained"]
    assert not a["feasible"] and any(c["code"] == "DEPENDENCY_ORDER" for c in a["constraints"])
    assert [(r["process_id"], r["start_minutes"], r["end_minutes"]) for r in b["rows"]] == [
        (2, 0, 60),
        (1, 60, 120),
    ]
    assert b["feasible"] and Decimal(b["total_cost"]) == 200 and b["utilization"]["r"] == 0.5


def test_parallelism_capacity_and_delay():
    s = example()
    s["dependencies"] = []
    result = schedule(s, "constrained")
    assert [(r["start_minutes"], r["end_minutes"]) for r in result["rows"]] == [(0, 60), (0, 60)]
    assert result["utilization"]["r"] == 1
    s["resources"]["r"]["capacity"] = 1
    result = schedule(s, "constrained")
    assert [(r["start_minutes"], r["end_minutes"]) for r in result["rows"]] == [(0, 60), (60, 120)]
    for t in range(result["makespan_minutes"]):
        assert (
            sum(
                r["assigned_resources"]["r"]
                for r in result["rows"]
                if r["start_minutes"] <= t < r["end_minutes"]
            )
            <= 1
        )
    s["resources"]["r"]["available"] = 150
    result = schedule(s, "constrained")
    assert result["rto_breaches"] == 2 and not result["feasible"]


def test_cycles_shortage_budget_and_impossible_target():
    s = example()
    s["dependencies"].append(
        {"from_type": "PROCESS", "from_id": 2, "to_type": "PROCESS", "to_id": 1}
    )
    result = schedule(s, "constrained")
    assert (
        not result["feasible"]
        and not result["rows"]
        and any(c["code"] == "DEPENDENCY_CYCLE" for c in result["constraints"])
    )
    for mutation in ["capacity", "budget", "rto"]:
        s = example()
        if mutation == "capacity":
            s["resources"]["r"]["capacity"] = 0
        if mutation == "budget":
            s["budget"] = "0"
        if mutation == "rto":
            s["tasks"][0]["target"] = 30
        assert not schedule(s, "constrained")["feasible"]


@pytest.mark.django_db
def test_remaining_crud_pagination_filters_and_normalized_seconds(data):
    client = data["client"]
    dept = create(client, "departments", {"name": "Sinov bo‘limi"})
    user = create(
        client,
        "users",
        {
            "username": "new-employee",
            "password": PASSWORD,
            "role": "PROCESS_OWNER",
            "department": dept["id"],
        },
    )
    assert (
        client.patch(
            f"/api/v1/users/{user['id']}/", {"first_name": "Yangilandi"}, format="json"
        ).status_code
        == 200
    )
    assert client.delete(f"/api/v1/users/{user['id']}/").status_code == 204
    assert User.objects.get(pk=user["id"]).is_active is False
    for resource, values, change in [
        ("departments", {"name": "Vaqtinchalik bo‘lim"}, {"head_name": "Yangi boshliq"}),
        ("services", {"name": "Vaqtinchalik xizmat"}, {"description": "Yangi tavsif"}),
        (
            "assets",
            {"name": "Vaqtinchalik aktiv", "type": "SERVER"},
            {"description": "Yangi tavsif"},
        ),
        ("resources", {"name": "Vaqtinchalik resurs", "capacity": 2}, {"capacity": 3}),
        (
            "strategies",
            {
                "process": data["process"].pk,
                "name": "Vaqtinchalik variant",
                "estimated_rto_hours": "1.50",
                "estimated_rpo_minutes": 30,
            },
            {"description": "Yangi tavsif"},
        ),
        (
            "risks",
            {
                "process": data["process"].pk,
                "threat": "Vaqtinchalik tahdid",
                "likelihood": 1,
                "impact": 2,
            },
            {"likelihood": 2},
        ),
        (
            "scenarios",
            {"title": "Vaqtinchalik ssenariy", "estimated_outage_hours": "1.50"},
            {"description": "Yangi tavsif"},
        ),
        (
            "plans",
            {"process": data["process"].pk, "title": "Qoralama reja"},
            {"scope": "Yangi qamrov"},
        ),
        (
            "bia",
            {"process": data["process"].pk, "rto_hours": "1.50", "mtpd_hours": "4.00"},
            {"comments": "Yangi izoh"},
        ),
    ]:
        obj = create(client, resource, values)
        url = f"/api/v1/{resource}/{obj['id']}/"
        assert client.get(url).status_code == 200
        response = client.patch(url, change, format="json")
        assert response.status_code == 200, response.data
        assert client.delete(url).status_code == 204
        assert client.get(url).status_code == 404
    for i in range(25):
        Asset.objects.create(organization=data["org"], name=f"Pagination-{i:02}", type="NETWORK")
    first = client.get("/api/v1/assets/?search=Pagination-&ordering=name&page_size=20").data
    second = client.get("/api/v1/assets/?search=Pagination-&ordering=name&page_size=20&page=2").data
    assert first["count"] == 25 and len(first["results"]) == 20 and len(second["results"]) == 5
    assert set(r["id"] for r in first["results"]).isdisjoint(r["id"] for r in second["results"])
    bia = BIAAssessment.objects.create(
        organization=data["org"],
        process=data["process"],
        rto_hours=Decimal("1.50"),
        revision=2,
        mtpd_hours=4,
        rpo_minutes=30,
    )
    assert bia.rto_seconds == 5400 and bia.mtpd_seconds == 14400 and bia.rpo_seconds == 1800


@pytest.mark.django_db
def test_login_throttle_and_secret_free_failed_audit():
    client = APIClient()
    for _ in range(10):
        assert (
            client.post(
                "/api/v1/auth/login", {"username": "missing", "password": PASSWORD}, format="json"
            ).status_code
            == 400
        )
    assert (
        client.post(
            "/api/v1/auth/login", {"username": "missing", "password": PASSWORD}, format="json"
        ).status_code
        == 429
    )
    assert AuditLog.objects.filter(action="LOGIN_FAILED").count() == 10
    assert PASSWORD not in json.dumps(list(AuditLog.objects.values("after_state")))


@pytest.mark.django_db
def test_last_administrator_cannot_be_disabled_or_demoted(data):
    url = f"/api/v1/users/{data['admin'].pk}/"
    assert data["client"].patch(url, {"role": "AUDITOR"}, format="json").status_code == 400
    assert data["client"].delete(url).status_code == 400


def test_multipage_unicode_pdf_and_docx_are_readable():
    from apps.plans.documents import export_pdf, export_docx

    values = plan_values(1)
    values.update(organization_name="O‘zbek o‘quv markazi", process_name="Ta’lim", status="DRAFT")
    values["steps"] = [
        {
            "sequence": i,
            "title": f"{i}. O‘quv xizmatini tiklash",
            "description": "Ma’lumotlar yaxlitligini tekshirish. " * 35,
            "expected_duration_minutes": 10,
            "assigned_name": "Mas’ul xodim",
        }
        for i in range(1, 35)
    ]
    pdf = PdfReader(io.BytesIO(export_pdf(values)))
    assert len(pdf.pages) > 1
    assert "34. O‘quv xizmatini tiklash" in "".join(page.extract_text() for page in pdf.pages)
    doc = Document(io.BytesIO(export_docx(values)))
    assert len(doc.tables[0].rows) == 35


@pytest.mark.django_db
def test_nested_steps_controls_exercises_and_simulation_history(data):
    c = data["client"]
    p = data["process"]
    org = data["org"]
    BIAAssessment.objects.create(
        organization=org, process=p, rto_hours=4, mtpd_hours=12, status="APPROVED"
    )
    plan = create(c, "plans", plan_values(p.pk))
    url = f"/api/v1/plans/{plan['id']}/"
    steps = [
        {
            "sequence": 1,
            "title": "Yangi bosqich",
            "assigned_to": data["admin"].pk,
            "expected_duration_minutes": 20,
        }
    ]
    assert c.patch(url, {"steps": steps}, format="json").status_code == 200
    assert c.get(url).data["steps"][0]["title"] == "Yangi bosqich"
    steps[0]["assigned_to"] = data["other_user"].pk
    assert c.patch(url, {"steps": steps}, format="json").status_code == 400
    assert c.post(url + "publish/").status_code == 200
    exercise = create(
        c,
        "exercises",
        {"plan": plan["id"], "exercise_date": "2026-10-21", "actual_recovery_hours": 1},
    )
    eurl = f"/api/v1/exercises/{exercise['id']}/"
    assert c.patch(eurl, {"findings": "Yangi kamchilik"}, format="json").status_code == 200
    assert c.delete(eurl).status_code == 204
    risk = create(
        c, "risks", {"process": p.pk, "threat": "Test tahdid", "likelihood": 3, "impact": 4}
    )
    control = create(c, "controls", {"risk": risk["id"], "title": "Test chora"})
    curl = f"/api/v1/controls/{control['id']}/"
    assert (
        c.patch(
            curl, {"status": "IMPLEMENTED", "effectiveness_evidence": "O‘quv dalili"}, format="json"
        ).status_code
        == 200
    )
    assert c.delete(curl).status_code == 204
    dep = create(
        c,
        "dependencies",
        {"from_type": "PROCESS", "from_id": p.pk, "to_type": "ASSET", "to_id": data["asset"].pk},
    )
    durl = f"/api/v1/dependencies/{dep['id']}/"
    assert c.patch(durl, {"notes": "Test aloqa"}, format="json").status_code == 200
    scenario = create(
        c, "scenarios", {"title": "History", "disrupted_asset_ids": [data["asset"].pk]}
    )
    result = c.post(f"/api/v1/scenarios/{scenario['id']}/simulate/")
    assert result.status_code == 200
    history = c.get("/api/v1/simulation-runs/").data["results"]
    assert history[0]["id"] == result.data["simulation_run_id"]
    assert history[0]["input_snapshot"]["scenario"]["disrupted_asset_ids"] == [data["asset"].pk]
    assert (
        c.patch(f"/api/v1/simulation-runs/{history[0]['id']}/", {}, format="json").status_code
        == 405
    )
    assert c.delete(durl).status_code == 204
    assert c.get("/api/v1/reports/risk-matrix").data["matrix"][2][3] == 1


@pytest.mark.django_db
@pytest.mark.parametrize("command", ["seed_demo_data", "provision_demo_users"])
def test_demo_commands_refuse_production(command, settings):
    from django.core.management import call_command
    from django.core.management.base import CommandError

    settings.DEBUG = False
    with pytest.raises(CommandError):
        call_command(command, verbosity=0)
    assert not Organization.objects.exists()


@pytest.mark.django_db
def test_explicit_demo_seed_never_resets_existing_password(monkeypatch, settings):
    from django.core.management import call_command

    settings.DEBUG = True
    monkeypatch.setenv("BCMS_DEMO_PASSWORD", PASSWORD)
    call_command("seed_demo_data", verbosity=0)
    user = User.objects.get(username="demo_admin")
    original = user.password
    monkeypatch.setenv("BCMS_DEMO_PASSWORD", "Another-Example!9126")
    call_command("provision_demo_users", verbosity=0)
    user.refresh_from_db()
    assert user.password == original and user.check_password(PASSWORD)
    assert user.organization.is_demo
