"""PDF and Excel export service."""

import io
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

from app.models.plan import ContinuityPlan
from app.models.process import BusinessProcess, Asset
from app.models.bia import BIAAssessment
from app.models.risk import Risk
from app.models.organization import Organization


def generate_plan_pdf(plan: ContinuityPlan) -> bytes:
    """Generate a PDF document for a continuity plan."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph(f"BCP/DRP Reja: {plan.title}", styles["Title"]))
    elements.append(Spacer(1, 10 * mm))

    # Meta info
    meta_data = [
        ["Jarayon:", plan.process.name if plan.process else "N/A"],
        ["Versiya:", plan.version],
        ["Holati:", plan.status],
        ["Qamrov:", plan.scope or "Ko'rsatilmagan"],
    ]
    if plan.approved_by:
        meta_data.append(["Tasdiqlovchi:", f"{plan.approved_by.first_name} {plan.approved_by.last_name}"])

    meta_table = Table(meta_data, colWidths=[40 * mm, 120 * mm])
    meta_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10 * mm))

    # Steps
    if plan.steps:
        elements.append(Paragraph("Tiklash Bosqichlari", styles["Heading2"]))
        elements.append(Spacer(1, 5 * mm))

        step_data = [["#", "Bosqich", "Mas'ul", "Vaqt (daqiqa)"]]
        for step in plan.steps:
            assigned = ""
            if step.assigned_to:
                assigned = f"{step.assigned_to.first_name} {step.assigned_to.last_name}"
            step_data.append([
                str(step.sequence),
                step.title,
                assigned or "Belgilanmagan",
                str(step.expected_duration_minutes),
            ])

        step_table = Table(step_data, colWidths=[15 * mm, 80 * mm, 45 * mm, 25 * mm])
        step_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4FA")]),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(step_table)

    # Footer
    elements.append(Spacer(1, 15 * mm))
    elements.append(Paragraph(
        "BCMS Designer — ISO 22301 Platform | Farg'ona Davlat Texnika Universiteti",
        styles["Normal"],
    ))

    doc.build(elements)
    return buffer.getvalue()


def generate_plan_excel(plan: ContinuityPlan) -> bytes:
    """Generate Excel workbook for a single continuity plan."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BCP Rejasi"

    # Header style
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A5F")
    border_thin = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD'),
    )

    # Title & metadata
    ws.cell(row=1, column=1, value=f"BCP/DRP Rejasi: {plan.title} (v{plan.version})").font = title_font
    ws.cell(row=3, column=1, value="Jarayon:").font = Font(bold=True)
    ws.cell(row=3, column=2, value=plan.process.name if plan.process else "N/A")
    ws.cell(row=4, column=1, value="Holati:").font = Font(bold=True)
    ws.cell(row=4, column=2, value=plan.status)
    ws.cell(row=5, column=1, value="Tasdiqlovchi:").font = Font(bold=True)
    ws.cell(row=5, column=2, value=f"{plan.approved_by.first_name} {plan.approved_by.last_name}" if plan.approved_by else "—")

    # Steps Table
    start_row = 7
    headers = ["#", "Tiklash bosqichi", "Mas'ul shaxs", "Kutilayotgan vaqt (daqiqa)"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 4] else "left")

    for idx, step in enumerate(plan.steps, 1):
        r = start_row + idx
        ws.cell(row=r, column=1, value=step.sequence).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=step.title)
        assigned = f"{step.assigned_to.first_name} {step.assigned_to.last_name}" if step.assigned_to else "Belgilanmagan"
        ws.cell(row=r, column=3, value=assigned)
        ws.cell(row=r, column=4, value=step.expected_duration_minutes).alignment = Alignment(horizontal="center")

        for c in range(1, 5):
            ws.cell(row=r, column=c).border = border_thin

    # Auto column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_organization_report_excel(db: Session, org_id: int) -> bytes:
    """Generate comprehensive multi-sheet Excel report for the organization."""
    wb = openpyxl.Workbook()
    org = db.query(Organization).filter(Organization.id == org_id).first()
    org_name = org.name if org else "Tashkilot"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD'),
    )

    # Sheet 1: Processes
    ws_proc = wb.active
    ws_proc.title = "Biznes Jarayonlar"
    proc_headers = ["ID", "Jarayon Nomi", "Kritiklik", "Holati", "Xizmat", "Mas'ul Egasi"]
    for c, h in enumerate(proc_headers, 1):
        cell = ws_proc.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill

    procs = db.query(BusinessProcess).filter(
        BusinessProcess.organization_id == org_id, BusinessProcess.is_deleted == False
    ).order_by(BusinessProcess.id).all()

    for r, p in enumerate(procs, 2):
        ws_proc.cell(row=r, column=1, value=p.id)
        ws_proc.cell(row=r, column=2, value=p.name)
        ws_proc.cell(row=r, column=3, value=p.criticality_display)
        ws_proc.cell(row=r, column=4, value=p.status)
        ws_proc.cell(row=r, column=5, value=p.service.name if p.service else "—")
        ws_proc.cell(row=r, column=6, value=f"{p.owner.first_name} {p.owner.last_name}".strip() if p.owner else "—")
        for col_idx in range(1, 7):
            ws_proc.cell(row=r, column=col_idx).border = border_thin

    # Sheet 2: BIA
    ws_bia = wb.create_sheet(title="BIA Baholash")
    bia_headers = ["ID", "Jarayon", "RTO (soat)", "MTPD (soat)", "RPO (daqiqa)", "Backup (daqiqa)", "RPO Bo'shliq", "Holati"]
    for c, h in enumerate(bia_headers, 1):
        cell = ws_bia.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill

    bias = db.query(BIAAssessment).filter(
        BIAAssessment.organization_id == org_id, BIAAssessment.is_deleted == False
    ).order_by(BIAAssessment.id).all()

    for r, b in enumerate(bias, 2):
        ws_bia.cell(row=r, column=1, value=b.id)
        ws_bia.cell(row=r, column=2, value=b.process.name if b.process else "—")
        ws_bia.cell(row=r, column=3, value=float(b.rto_hours))
        ws_bia.cell(row=r, column=4, value=float(b.mtpd_hours))
        ws_bia.cell(row=r, column=5, value=b.rpo_minutes)
        ws_bia.cell(row=r, column=6, value=b.backup_interval_minutes)
        ws_bia.cell(row=r, column=7, value="GAP MAVJUD" if b.has_rpo_gap else "QONIQARLI")
        ws_bia.cell(row=r, column=8, value=b.status)
        for col_idx in range(1, 9):
            ws_bia.cell(row=r, column=col_idx).border = border_thin

    # Sheet 3: Risks
    ws_risk = wb.create_sheet(title="Risk Reestri")
    risk_headers = ["ID", "Jarayon", "Tahdid", "Aktiv", "Ehtimollik", "Ta'sir", "Risk Balli", "Qoldiq Ball", "Klass", "Holati"]
    for c, h in enumerate(risk_headers, 1):
        cell = ws_risk.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill

    risks = db.query(Risk).filter(
        Risk.organization_id == org_id, Risk.is_deleted == False
    ).order_by(Risk.id).all()

    for r, rk in enumerate(risks, 2):
        ws_risk.cell(row=r, column=1, value=rk.id)
        ws_risk.cell(row=r, column=2, value=rk.process.name if rk.process else "—")
        ws_risk.cell(row=r, column=3, value=rk.threat)
        ws_risk.cell(row=r, column=4, value=rk.asset.name if rk.asset else "—")
        ws_risk.cell(row=r, column=5, value=rk.likelihood)
        ws_risk.cell(row=r, column=6, value=rk.impact)
        ws_risk.cell(row=r, column=7, value=rk.risk_score)
        ws_risk.cell(row=r, column=8, value=rk.residual_risk_score or "—")
        ws_risk.cell(row=r, column=9, value=rk.get_risk_class().upper())
        ws_risk.cell(row=r, column=10, value=rk.status)
        for col_idx in range(1, 11):
            ws_risk.cell(row=r, column=col_idx).border = border_thin

    # Format column widths across all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
